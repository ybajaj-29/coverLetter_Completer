import docx
from fpdf import FPDF

def read_word_file(filename):
    doc = docx.Document(filename)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

def replace_placeholders(text, placeholders):
    for placeholder in placeholders:
        replacement = input(f"Enter a value for {placeholder}: ")
        text = text.replace(placeholder, replacement)
    return text

def save_to_pdf(text, filename):
    pdf = FPDF()
    pdf.add_page()
    
    # Add the Calibri font to the FPDF instance
    pdf.add_font('Calibri', '', 'calibri.ttf', uni=True)
    pdf.set_font("Calibri", size=11)  
    
    pdf.multi_cell(0, 5, text) # the second parameter controls the line spacing, where 5 was found to be the most legible
    pdf.output(filename)

if __name__ == "__main__":
    cover_letter_template = read_word_file("coverLetter_template.docx")
    
    placeholders = [
        "[Today’s Date]",
        "[Company's Address]",
        "[City, State, ZIP Code]",
        "[Position Name]",
        "[Company Name]",
        "[specific aspect(s) about the company or role]"
    ]

    updated_cover_letter = replace_placeholders(cover_letter_template, placeholders)
    save_to_pdf(updated_cover_letter, "C:\\Users\\Yasha\\funWithPython\\coverLetter_Completer\\Cover Letter.pdf")
    print("Cover letter saved to 'Cover Letter.pdf'")